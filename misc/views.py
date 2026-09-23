from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment
from students.models import Student
from accounts.models import Group, Section
from .models import StudentExamInfo

# Group-specific IPE subjects (s3..s6 slots)
GROUP_IPE_SUBJECTS = {
    'MPIC': ['Physics', 'Chemistry', 'Maths',       ''],
    'BPIC': ['Biology', 'Zoology',   'Chemistry',   ''],
    'CEC':  ['Civics',  'Economics', 'Commerce',    'Sanskrit'],
    'MEC':  ['Mathematics', 'Economics', 'Commerce','Sanskrit'],
}
DEFAULT_IPE_SUBJECTS = ['Sub 3', 'Sub 4', 'Sub 5', 'Sub 6']

IPE_MARK_FIELDS_1ST = ['ipe_1st_telugu','ipe_1st_english','ipe_1st_s3','ipe_1st_s4','ipe_1st_s5','ipe_1st_s6','ipe_1st_total']
IPE_MARK_FIELDS_2ND = ['ipe_2nd_telugu','ipe_2nd_english','ipe_2nd_s3','ipe_2nd_s4','ipe_2nd_s5','ipe_2nd_s6','ipe_2nd_total']


def _group_subjects(group_code):
    return GROUP_IPE_SUBJECTS.get(group_code.upper(), DEFAULT_IPE_SUBJECTS)


def _ensure_exam_info(students):
    ids = [s.id for s in students]
    existing = set(StudentExamInfo.objects.filter(student_id__in=ids).values_list('student_id', flat=True))
    to_create = [StudentExamInfo(student_id=sid) for sid in ids if sid not in existing]
    if to_create:
        StudentExamInfo.objects.bulk_create(to_create)


@login_required
def misc_home(request):
    return redirect('misc_sheet')


@login_required
def misc_sheet(request):
    groups   = Group.objects.prefetch_related('sections').order_by('name')
    sections = Section.objects.select_related('group').order_by('group__name', 'year', 'name')

    sel_groups   = request.GET.getlist('groups')
    sel_sections = request.GET.getlist('sections')

    students = Student.objects.filter(is_active=True).select_related(
        'section__group', 'exam_info'
    ).order_by('section__group__name', 'section__year', 'name')

    if sel_groups:
        students = students.filter(section__group_id__in=sel_groups)
    if sel_sections:
        students = students.filter(section_id__in=sel_sections)

    students = list(students)
    _ensure_exam_info(students)
    # re-fetch with exam_info
    ids = [s.id for s in students]
    students = list(
        Student.objects.filter(id__in=ids).select_related('section__group', 'exam_info')
        .order_by('section__group__name', 'section__year', 'name')
    )

    # Build per-student group subjects mapping for IPE columns
    def student_subs(s):
        code = s.section.group.code if s.section and s.section.group else ''
        return _group_subjects(code)

    tab = request.GET.get('tab', 'ipe')

    return render(request, 'misc/exam_sheet.html', {
        'students': students,
        'groups': groups,
        'sections': sections,
        'sel_groups': sel_groups,
        'sel_sections': sel_sections,
        'tab': tab,
        'student_subs': {s.id: student_subs(s) for s in students},
    })


@login_required
def misc_save(request):
    if request.method != 'POST':
        return redirect('misc_sheet')

    tab = request.POST.get('tab', 'ipe')
    ids = request.POST.getlist('student_ids')

    for sid in ids:
        try:
            info = StudentExamInfo.objects.get(student_id=sid)
        except StudentExamInfo.DoesNotExist:
            info = StudentExamInfo(student_id=sid)

        student = info.student

        def val(key):
            return request.POST.get(f'{key}_{sid}', '').strip()

        def dec(key):
            v = val(key)
            return v if v else None

        def dt(key):
            v = val(key)
            return v if v else None

        if tab == 'ipe':
            # student DOB & IPE hall ticket
            dob = dt('dob')
            if dob:
                student.date_of_birth = dob
                student.save(update_fields=['date_of_birth'])
            info.ipe_hall_ticket = val('ipe_ht')
            for f in ['ipe_1st_telugu','ipe_1st_english','ipe_1st_s3','ipe_1st_s4',
                      'ipe_1st_s5','ipe_1st_s6','ipe_1st_total',
                      'ipe_2nd_telugu','ipe_2nd_english','ipe_2nd_s3','ipe_2nd_s4',
                      'ipe_2nd_s5','ipe_2nd_s6','ipe_2nd_total']:
                setattr(info, f, dec(f))

        elif tab == 'eamcet':
            dob = dt('dob')
            if dob:
                student.date_of_birth = dob
                student.save(update_fields=['date_of_birth'])
            info.eamcet_hall_ticket  = val('eamcet_ht')
            info.eamcet_password     = val('eamcet_pwd')
            info.eamcet_date_of_exam = dt('eamcet_date')
            info.eamcet_marks        = dec('eamcet_marks')
            info.eamcet_rank         = val('eamcet_rank')

        elif tab == 'jee':
            dob = dt('dob')
            if dob:
                student.date_of_birth = dob
                student.save(update_fields=['date_of_birth'])
            info.jee_hall_ticket  = val('jee_ht')
            info.jee_password     = val('jee_pwd')
            info.jee_date_of_exam = dt('jee_date')

        info.save()

    messages.success(request, 'Data saved successfully.')
    # Preserve filters + tab
    qs = request.POST.get('next_qs', '')
    return redirect(f'/misc/?{qs}' if qs else 'misc_sheet')


@login_required
def misc_export(request):
    from openpyxl.styles import Font, PatternFill, Alignment
    tab      = request.GET.get('tab', 'ipe')
    sel_grps = request.GET.getlist('groups')
    sel_secs = request.GET.getlist('sections')

    students = Student.objects.filter(is_active=True).select_related('section__group', 'exam_info').order_by('section__group__name', 'section__year', 'name')
    if sel_grps:
        students = students.filter(section__group_id__in=sel_grps)
    if sel_secs:
        students = students.filter(section_id__in=sel_secs)
    students = list(students)
    _ensure_exam_info(students)
    students = list(Student.objects.filter(id__in=[s.id for s in students]).select_related('section__group', 'exam_info').order_by('section__group__name', 'section__year', 'name'))

    wb = openpyxl.Workbook()
    ws = wb.active

    hf   = Font(bold=True, color='FFFFFF')
    hfil = PatternFill(fill_type='solid', fgColor='1a56db')
    ha   = Alignment(horizontal='center', vertical='center', wrap_text=True)

    def hdr(ws, headers, widths):
        for i, h in enumerate(headers, 1):
            c = ws.cell(1, i, h); c.font = hf; c.fill = hfil; c.alignment = ha
        ws.row_dimensions[1].height = 24
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'A2'

    if tab == 'ipe':
        ws.title = 'IPE'
        headers = ['#','Adm No','Name','Sec','Phone','DOB','Parent Name','IPE Hall Ticket',
                   '10th Total',
                   '1Y Telugu','1Y English','1Y S3','1Y S4','1Y S5','1Y S6','1Y Total',
                   '2Y Telugu','2Y English','2Y S3','2Y S4','2Y S5','2Y S6','2Y Total']
        widths  = [5,13,26,8,13,13,20,18,10,
                   10,10,12,12,12,12,10,
                   10,10,12,12,12,12,10]
        hdr(ws, headers, widths)
        for r, s in enumerate(students, 2):
            inf = getattr(s, 'exam_info', None)
            def m(f): return float(getattr(inf, f)) if inf and getattr(inf,f) is not None else ''
            ws.cell(r,1,r-1); ws.cell(r,2,s.admission_number); ws.cell(r,3,s.name)
            ws.cell(r,4,s.section.name if s.section else '')
            ws.cell(r,5,s.mobile); ws.cell(r,6,str(s.date_of_birth or ''))
            ws.cell(r,7,s.father_name)
            ws.cell(r,8,inf.ipe_hall_ticket if inf else '')
            ws.cell(r,9,float(s.marks_total) if s.marks_total else '')
            for i,f in enumerate(['ipe_1st_telugu','ipe_1st_english','ipe_1st_s3','ipe_1st_s4','ipe_1st_s5','ipe_1st_s6','ipe_1st_total'], 10):
                ws.cell(r, i, m(f))
            for i,f in enumerate(['ipe_2nd_telugu','ipe_2nd_english','ipe_2nd_s3','ipe_2nd_s4','ipe_2nd_s5','ipe_2nd_s6','ipe_2nd_total'], 17):
                ws.cell(r, i, m(f))

    elif tab == 'eamcet':
        ws.title = 'EAMCET'
        headers = ['#','Adm No','Name','Sec','Phone','DOB','Parents No',
                   'Hall Ticket','Password','Date of Exam','Marks','Rank']
        widths  = [5,13,26,8,13,13,13,20,20,14,12,14]
        hdr(ws, headers, widths)
        for r, s in enumerate(students, 2):
            inf = getattr(s, 'exam_info', None)
            ws.cell(r,1,r-1); ws.cell(r,2,s.admission_number); ws.cell(r,3,s.name)
            ws.cell(r,4,s.section.name if s.section else '')
            ws.cell(r,5,s.mobile); ws.cell(r,6,str(s.date_of_birth or ''))
            ws.cell(r,7,s.mobile)
            ws.cell(r,8,inf.eamcet_hall_ticket if inf else '')
            ws.cell(r,9,'••••••' if (inf and inf.eamcet_password) else '')
            ws.cell(r,10,str(inf.eamcet_date_of_exam or '') if inf else '')
            ws.cell(r,11,float(inf.eamcet_marks) if inf and inf.eamcet_marks else '')
            ws.cell(r,12,inf.eamcet_rank if inf else '')

    else:  # jee
        ws.title = 'JEE Main'
        headers = ['#','Adm No','Name','Sec','DOB','Parents No',
                   'Hall Ticket','Password','Date of Exam']
        widths  = [5,13,26,8,13,13,20,20,14]
        hdr(ws, headers, widths)
        for r, s in enumerate(students, 2):
            inf = getattr(s, 'exam_info', None)
            ws.cell(r,1,r-1); ws.cell(r,2,s.admission_number); ws.cell(r,3,s.name)
            ws.cell(r,4,s.section.name if s.section else '')
            ws.cell(r,5,str(s.date_of_birth or ''))
            ws.cell(r,6,s.mobile)
            ws.cell(r,7,inf.jee_hall_ticket if inf else '')
            ws.cell(r,8,'••••••' if (inf and inf.jee_password) else '')
            ws.cell(r,9,str(inf.jee_date_of_exam or '') if inf else '')

    tab_name = {'ipe': 'IPE', 'eamcet': 'EAMCET', 'jee': 'JEE_Main'}.get(tab, tab)
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename=misc_{tab_name}.xlsx'
    wb.save(response)
    return response
