from django.shortcuts import render, redirect, get_object_or_404
from django.conf import settings
from django.http import HttpResponse
from django.contrib import messages
from django.db.models import Q, Sum
from django.utils import timezone
import uuid
from .models import StudentFee, FeePayment, FeeType, StudentFeeCharge
from .receipt_pdf import generate_receipt_pdf, save_receipt_pdf
from students.models import Student
from accounts.models import AcademicYear
from accounts.decorators import admin_accounts_required, all_roles_required
from whatsapp.services import send_whatsapp_media

@admin_accounts_required
def fee_type_manage(request):
    active_year = AcademicYear.objects.filter(is_active=True).first()
    if not active_year:
        messages.error(request, "No active academic year found.")
        return redirect('fee_list')

    if request.method == 'POST':
        name = request.POST.get('name', '').strip()

        if name:
            fee_type, created = FeeType.objects.get_or_create(name=name, academic_year=active_year)
            if created:
                messages.success(request, f"Fee type '{name}' created. You can now assign it to sections.")
            else:
                messages.warning(request, f"Fee type '{name}' already exists for this year.")
        return redirect('fee_type_manage')

    fee_types = FeeType.objects.filter(academic_year=active_year).order_by('-created_at')
    
    # Calculate stats per fee type
    for ft in fee_types:
        charges = StudentFeeCharge.objects.filter(fee_type=ft).prefetch_related('payments')
        ft.total_assigned = sum(c.amount_assigned for c in charges)
        ft.total_paid = sum(c.total_paid for c in charges)
        ft.total_pending = ft.total_assigned - ft.total_paid

    # Calculate overall Tuition Fee stats for active academic year
    tuition_fees = StudentFee.objects.filter(academic_year=active_year).prefetch_related('payments')
    tuition_stats = {
        'name': 'Tuition Fee (Core Fee)',
        'total_assigned': sum(sf.total_fee for sf in tuition_fees),
        'total_paid': sum(sf.total_paid for sf in tuition_fees),
        'total_pending': sum(sf.total_pending for sf in tuition_fees),
    }

    return render(request, 'fees/manage_types.html', {
        'active_year': active_year,
        'fee_types': fee_types,
        'tuition_stats': tuition_stats,
    })


@admin_accounts_required
def fee_type_edit(request, pk):
    active_year = AcademicYear.objects.filter(is_active=True).first()
    fee_type = get_object_or_404(FeeType, pk=pk, academic_year=active_year)

    if request.method == 'POST':
        new_name = request.POST.get('name', '').strip()
        if new_name:
            exists = FeeType.objects.filter(name__iexact=new_name, academic_year=active_year).exclude(pk=pk).exists()
            if exists:
                messages.error(request, f"A fee type named '{new_name}' already exists.")
            else:
                old_name = fee_type.name
                fee_type.name = new_name
                fee_type.save()
                messages.success(request, f"Fee type '{old_name}' renamed to '{new_name}'.")
        return redirect('fee_type_manage')

    return redirect('fee_type_manage')


@admin_accounts_required
def fee_type_assign(request, pk):
    active_year = AcademicYear.objects.filter(is_active=True).first()
    fee_type = get_object_or_404(FeeType, pk=pk, academic_year=active_year)

    from accounts.models import Section
    sections = list(Section.objects.select_related('group').all().order_by('group__name', 'year', 'name'))

    charges = StudentFeeCharge.objects.filter(fee_type=fee_type, student__is_active=True).select_related('student')
    section_charges = {}
    for c in charges:
        sec_id = c.student.section_id
        if sec_id:
            section_charges.setdefault(sec_id, []).append(c)

    for s in sections:
        c_list = section_charges.get(s.id, [])
        if c_list:
            s.is_assigned = True
            s.assigned_amount = c_list[0].amount_assigned
            s.assigned_count = len(c_list)
        else:
            s.is_assigned = False
            s.assigned_amount = 0
            s.assigned_count = 0

    if request.method == 'POST':
        section_ids = request.POST.getlist('sections')
        amount_str = request.POST.get('amount', '0').strip()
        try:
            amount = float(amount_str)
        except ValueError:
            amount = 0

        if not section_ids:
            messages.error(request, 'Please select at least one section.')
            return redirect('fee_type_assign', pk=pk)

        students = Student.objects.filter(is_active=True, academic_year=active_year, section_id__in=section_ids)
        updated_count = 0
        for student in students:
            charge, _ = StudentFeeCharge.objects.get_or_create(
                student=student, 
                fee_type=fee_type,
                defaults={'amount_assigned': amount}
            )
            charge.amount_assigned = amount
            charge.save()
            updated_count += 1
        
        messages.success(request, f"Fee '{fee_type.name}' updated to ₹{amount:.0f} for {updated_count} students.")
        return redirect('fee_type_manage')

    return render(request, 'fees/assign_type.html', {
        'fee_type': fee_type,
        'sections': sections,
    })


@admin_accounts_required
def fee_type_unassign_section(request, pk, section_id):
    active_year = AcademicYear.objects.filter(is_active=True).first()
    fee_type = get_object_or_404(FeeType, pk=pk, academic_year=active_year)
    from accounts.models import Section
    section = get_object_or_404(Section, pk=section_id)

    charges = StudentFeeCharge.objects.filter(fee_type=fee_type, student__section=section).prefetch_related('payments')
    paid_charges = [c for c in charges if c.total_paid > 0]

    if paid_charges:
        messages.error(request, f"Cannot unassign '{fee_type.name}' from {section} because payments have already been collected for {len(paid_charges)} student(s).")
    else:
        deleted_count, _ = charges.delete()
        messages.success(request, f"Unassigned '{fee_type.name}' from {section} ({deleted_count} student record(s) removed).")

    next_url = request.POST.get('next') or request.GET.get('next')
    if next_url:
        return redirect(next_url)
    return redirect('fee_type_assign', pk=pk)


@admin_accounts_required
def fee_type_assign_individual(request, pk):
    """Assign a fee type to students with a DIFFERENT hand-typed amount per
    student, instead of one bulk amount for a whole section. Also allows recording
    partial/installment payments directly from this screen."""
    active_year = AcademicYear.objects.filter(is_active=True).first()
    fee_type = get_object_or_404(FeeType, pk=pk, academic_year=active_year)

    from accounts.models import Section

    if request.method == 'POST':
        year_filter = request.POST.get('year_filter', '')
        section_filter = request.POST.get('section_filter', '')
        status_filter = request.POST.get('status_filter', '')
        ids = request.POST.getlist('student_ids')
        updated = 0
        payments_collected = 0

        for sid in ids:
            key = f'amount_{sid}'
            pay_key = f'pay_{sid}'
            raw_amt = request.POST.get(key, '').strip() if key in request.POST else ''
            raw_pay = request.POST.get(pay_key, '').strip() if pay_key in request.POST else ''

            student = Student.objects.filter(pk=sid).first()
            if not student:
                continue

            if raw_amt != '':
                try:
                    amount = float(raw_amt)
                    charge, _ = StudentFeeCharge.objects.get_or_create(
                        student=student, fee_type=fee_type, defaults={'amount_assigned': amount}
                    )
                    charge.amount_assigned = amount
                    charge.save()
                    updated += 1
                except ValueError:
                    pass

            if raw_pay != '':
                try:
                    pay_amount = float(raw_pay)
                    if pay_amount > 0:
                        charge, _ = StudentFeeCharge.objects.get_or_create(
                            student=student, fee_type=fee_type, defaults={'amount_assigned': 0}
                        )
                        receipt_number = 'RCP' + str(uuid.uuid4())[:8].upper()
                        FeePayment.objects.create(
                            fee_charge=charge,
                            amount=pay_amount,
                            payment_date=timezone.localdate(),
                            payment_mode='cash',
                            receipt_number=receipt_number,
                            collected_by=request.user.get_full_name() or request.user.username,
                            remarks=f"Payment for {fee_type.name} via Manage Fee",
                        )
                        payments_collected += 1
                except ValueError:
                    pass

        msg = f"'{fee_type.name}' updated for {updated} student(s)."
        if payments_collected > 0:
            msg += f" Recorded {payments_collected} partial payment(s)."
        messages.success(request, msg)
        return redirect(f"/fees/types/{pk}/assign-individual/?year={year_filter}&section={section_filter}&status={status_filter}")

    year_filter = request.GET.get('year', '')
    section_filter = request.GET.get('section', '')
    status_filter = request.GET.get('status', '').lower().strip()

    students = Student.objects.filter(is_active=True, academic_year=active_year).select_related('section__group')
    if year_filter:
        students = students.filter(section__year=year_filter)
    if section_filter:
        students = students.filter(section_id=section_filter)
    students = students.order_by('section__group__name', 'section__year', 'section__name', 'name')

    existing_charges = {
        c.student_id: c
        for c in StudentFeeCharge.objects.filter(fee_type=fee_type).prefetch_related('payments')
    }
    student_rows = [{'student': s, 'charge': existing_charges.get(s.pk)} for s in students]

    if status_filter == 'pending':
        student_rows = [r for r in student_rows if r['charge'] and r['charge'].status.lower() == 'pending']
    elif status_filter == 'partial':
        student_rows = [r for r in student_rows if r['charge'] and r['charge'].status.lower() == 'partial']
    elif status_filter == 'paid':
        student_rows = [r for r in student_rows if r['charge'] and r['charge'].status.lower() == 'paid']

    all_fee_types = FeeType.objects.filter(academic_year=active_year).order_by('name')
    sections = Section.objects.select_related('group').all().order_by('group__name', 'year', 'name')

    return render(request, 'fees/assign_type_individual.html', {
        'fee_type': fee_type,
        'all_fee_types': all_fee_types,
        'student_rows': student_rows,
        'sections': sections,
        'year_filter': year_filter,
        'section_filter': section_filter,
        'status_filter': status_filter,
    })


@admin_accounts_required
def fee_charge_delete(request, pk, charge_id):
    active_year = AcademicYear.objects.filter(is_active=True).first()
    fee_type = get_object_or_404(FeeType, pk=pk, academic_year=active_year)
    charge = get_object_or_404(StudentFeeCharge, pk=charge_id, fee_type=fee_type)

    if charge.total_paid > 0:
        messages.error(request, f"Cannot delete '{fee_type.name}' charge for {charge.student.name} because a payment of ₹{charge.total_paid:,.0f} has already been recorded.")
    else:
        student_name = charge.student.name
        charge.delete()
        messages.success(request, f"Removed '{fee_type.name}' charge for {student_name}.")

    year_filter = request.GET.get('year', '')
    section_filter = request.GET.get('section', '')
    status_filter = request.GET.get('status', '')
    return redirect(f"/fees/types/{pk}/assign-individual/?year={year_filter}&section={section_filter}&status={status_filter}")


@admin_accounts_required
def fee_type_delete(request, pk):
    active_year = AcademicYear.objects.filter(is_active=True).first()
    fee_type = get_object_or_404(FeeType, pk=pk, academic_year=active_year)
    
    if request.method == 'POST':
        name = fee_type.name
        fee_type.delete()  # Hard delete as requested
        messages.success(request, f"Fee type '{name}' has been completely deleted along with all its records.")
        return redirect('fee_type_manage')
        
    return render(request, 'fees/delete_type_confirm.html', {'fee_type': fee_type})


@all_roles_required
def fee_list(request):
    q = request.GET.get('q', '')
    status_filter = request.GET.get('status', 'pending').lower().strip()
    fee_type_filter = request.GET.get('fee_type', '').strip()
    section_filter = request.GET.get('section', '').strip()
    active_year = AcademicYear.objects.filter(is_active=True).first()

    students = Student.objects.filter(is_active=True).select_related('section__group')
    if section_filter:
        students = students.filter(section_id=section_filter)
    if q:
        students = students.filter(
            Q(name__icontains=q) | Q(admission_number__icontains=q)
        )
    student_list = list(students)
    student_ids = [s.pk for s in student_list]

    fee_types = list(FeeType.objects.filter(academic_year=active_year).order_by('created_at'))
    fee_data = []

    total_collected_tuition = 0
    total_pending_tuition = 0
    total_fees_tuition = 0

    def classify_status(total_fee_or_assigned, total_paid, total_pending):
        if total_fee_or_assigned <= 0:
            return 'not_started' if total_paid <= 0 else 'partial'
        if total_paid <= 0:
            return 'not_started'
        if total_pending <= 0:
            return 'paid'
        return 'partial'

    if active_year and student_ids:
        existing_fees = {
            sf.student_id: sf
            for sf in StudentFee.objects.filter(
                student_id__in=student_ids, academic_year=active_year
            ).prefetch_related('payments')
        }

        missing_ids = [pk for pk in student_ids if pk not in existing_fees]
        if missing_ids:
            StudentFee.objects.bulk_create(
                [StudentFee(student_id=pk, academic_year=active_year, total_fee=0) for pk in missing_ids],
                ignore_conflicts=True,
            )
            for sf in StudentFee.objects.filter(
                student_id__in=missing_ids, academic_year=active_year
            ).prefetch_related('payments'):
                existing_fees[sf.student_id] = sf

        charges_by_student = {}
        if fee_types:
            for c in StudentFeeCharge.objects.filter(
                student_id__in=student_ids, fee_type__academic_year=active_year
            ).select_related('fee_type').prefetch_related('payments'):
                charges_by_student.setdefault(c.student_id, {})[c.fee_type_id] = c

        for student in student_list:
            sf = existing_fees.get(student.pk)
            if not sf:
                continue
            sf.student = student

            charges_dict = charges_by_student.get(student.pk, {})
            ordered_charges = [charges_dict.get(ft.id) for ft in fee_types]

            # Compute scoped status
            if fee_type_filter == 'tuition':
                row_status = classify_status(sf.total_fee, sf.total_paid, sf.total_pending)
            elif fee_type_filter.isdigit():
                ft_id = int(fee_type_filter)
                target_charge = charges_dict.get(ft_id)
                if target_charge:
                    row_status = classify_status(target_charge.amount_assigned, target_charge.total_paid, target_charge.total_pending)
                else:
                    row_status = classify_status(0, 0, 0)
            else:
                combined_assigned = sf.total_fee + sum(c.amount_assigned for c in ordered_charges if c)
                combined_paid = sf.total_paid + sum(c.total_paid for c in ordered_charges if c)
                combined_pending = sf.total_pending + sum(c.total_pending for c in ordered_charges if c)
                row_status = classify_status(combined_assigned, combined_paid, combined_pending)

            # Filter against status_filter
            if status_filter == 'not_started' and row_status != 'not_started':
                continue
            elif status_filter == 'pending' and row_status not in ('partial', 'not_started'):
                continue
            elif status_filter == 'partial' and row_status != 'partial':
                continue
            elif status_filter == 'paid' and row_status != 'paid':
                continue

            fee_data.append({
                'student_fee': sf,
                'charges': ordered_charges,
                'overall_status': row_status,
            })

            total_collected_tuition += sf.total_paid
            total_pending_tuition += sf.total_pending
            total_fees_tuition += sf.total_fee

    total_charges_fees = sum(
        c.amount_assigned for row in fee_data for c in row['charges'] if c
    )
    total_charges_paid = sum(
        c.total_paid for row in fee_data for c in row['charges'] if c
    )
    total_charges_pending = sum(
        c.total_pending for row in fee_data for c in row['charges'] if c
    )

    total_fees = total_fees_tuition + total_charges_fees
    total_collected = total_collected_tuition + total_charges_paid
    total_pending = total_pending_tuition + total_charges_pending

    from accounts.models import Section
    sections = Section.objects.select_related('group').all().order_by('group__name', 'year', 'name')

    return render(request, 'fees/list.html', {
        'fee_data': fee_data,
        'fee_types': fee_types,
        'sections': sections,
        'q': q,
        'status_filter': status_filter,
        'fee_type_filter': fee_type_filter,
        'section_filter': section_filter,
        'total_fees': total_fees,
        'total_collected': total_collected,
        'total_pending': total_pending,
    })


@all_roles_required
def fee_export(request):
    """Export fees filtered by section, fee type, status, and search query to PDF or Excel."""
    q = request.GET.get('q', '')
    status_filter = request.GET.get('status', 'pending').lower().strip()
    fee_type_filter = request.GET.get('fee_type', '').strip()
    section_filter = request.GET.get('section', '').strip()
    fmt = request.GET.get('fmt', 'excel')
    active_year = AcademicYear.objects.filter(is_active=True).first()

    students = Student.objects.filter(is_active=True).select_related('section__group')
    if section_filter:
        students = students.filter(section_id=section_filter)
    if q:
        students = students.filter(
            Q(name__icontains=q) | Q(admission_number__icontains=q)
        )
    student_list = list(students)
    student_ids = [s.pk for s in student_list]

    fee_types = list(FeeType.objects.filter(academic_year=active_year).order_by('created_at'))

    fee_rows = []
    if active_year and student_ids:
        existing_fees = {
            sf.student_id: sf
            for sf in StudentFee.objects.filter(
                student_id__in=student_ids, academic_year=active_year
            ).prefetch_related('payments')
        }

        charges_by_student = {}
        if fee_types:
            for c in StudentFeeCharge.objects.filter(
                student_id__in=student_ids, fee_type__academic_year=active_year
            ).select_related('fee_type').prefetch_related('payments'):
                charges_by_student.setdefault(c.student_id, {})[c.fee_type_id] = c

        for student in student_list:
            sf = existing_fees.get(student.pk)
            if not sf:
                continue

            charges_dict = charges_by_student.get(student.pk, {})
            ordered_charges = [charges_dict.get(ft.id) for ft in fee_types]

            if fee_type_filter == 'tuition':
                if status_filter == 'pending' and sf.total_pending <= 0:
                    continue
                elif status_filter == 'partial' and sf.status != 'Partial':
                    continue
                elif status_filter == 'paid' and sf.status != 'Paid':
                    continue
            elif fee_type_filter.isdigit():
                ft_id = int(fee_type_filter)
                target_charge = charges_dict.get(ft_id)
                if status_filter == 'pending' and (not target_charge or target_charge.total_pending <= 0):
                    continue
                elif status_filter == 'partial' and (not target_charge or target_charge.status != 'Partial'):
                    continue
                elif status_filter == 'paid' and (not target_charge or target_charge.status != 'Paid'):
                    continue
            else:
                has_pending = (sf.total_pending > 0) or any(c and c.total_pending > 0 for c in ordered_charges if c)
                has_partial = (sf.status == 'Partial') or any(c and c.status == 'Partial' for c in ordered_charges if c)
                is_all_paid = (sf.status == 'Paid') and all((not c) or (c.status == 'Paid') for c in ordered_charges if c)

                if status_filter == 'pending' and not has_pending:
                    continue
                elif status_filter == 'partial' and not has_partial:
                    continue
                elif status_filter == 'paid' and not is_all_paid:
                    continue

            if fee_type_filter == 'tuition':
                row_total = float(sf.total_fee)
                row_paid = float(sf.total_paid)
                row_pending = float(sf.total_pending)
            elif fee_type_filter.isdigit():
                target_charge = charges_dict.get(int(fee_type_filter))
                row_total = float(target_charge.amount_assigned) if target_charge else 0
                row_paid = float(target_charge.total_paid) if target_charge else 0
                row_pending = float(target_charge.total_pending) if target_charge else 0
            else:
                row_total = float(sf.total_fee) + sum(float(c.amount_assigned) for c in ordered_charges if c)
                row_paid = float(sf.total_paid) + sum(float(c.total_paid) for c in ordered_charges if c)
                row_pending = float(sf.total_pending) + sum(float(c.total_pending) for c in ordered_charges if c)

            fee_rows.append({
                'name':    student.name,
                'section': str(student.section) if student.section else '—',
                'year':    student.section.get_year_display() if student.section else '—',
                'total':   row_total,
                'paid':    row_paid,
                'pending': row_pending,
                'phone':   student.mobile or '—',
            })

    if fmt == 'pdf':
        return _fee_export_pdf(fee_rows, active_year)
    return _fee_export_excel(fee_rows, active_year)


@admin_accounts_required
def fee_set(request, pk):
    student = get_object_or_404(Student, pk=pk)
    active_year = AcademicYear.objects.filter(is_active=True).first()
    student_fee, _ = StudentFee.objects.get_or_create(
        student=student, academic_year=active_year, defaults={'total_fee': 0}
    )
    if request.method == 'POST':
        raw_fee = request.POST.get('total_fee', '0').strip()
        try:
            student_fee.total_fee = float(raw_fee) if raw_fee else 0
        except ValueError:
            student_fee.total_fee = 0
        student_fee.save()
        messages.success(request, f'Total fee set to ₹{student_fee.total_fee} for {student.name}.')
        # FIX: clean redirect logic — check next first, else go to fee_detail
        next_url = request.POST.get('next', '').strip()
        if next_url:
            return redirect(next_url)
        return redirect('fee_detail', pk=pk)
    return render(request, 'fees/set_fee.html', {
        'student': student, 'student_fee': student_fee, 'active_year': active_year,
    })


@admin_accounts_required
def fee_set_bulk(request):
    active_year = AcademicYear.objects.filter(is_active=True).first()
    if not active_year:
        messages.error(request, "No active academic year found.")
        return redirect('fee_list')

    from accounts.models import Section
    sections = Section.objects.select_related('group').all().order_by('group__name', 'year', 'name')

    if request.method == 'POST':
        section_ids = request.POST.getlist('sections')
        amount_str = request.POST.get('amount', '0').strip()
        try:
            amount = float(amount_str)
        except ValueError:
            amount = 0

        if not section_ids:
            messages.error(request, 'Please select at least one section.')
            return redirect('fee_set_bulk')

        students = Student.objects.filter(is_active=True, academic_year=active_year, section_id__in=section_ids)
        updated_count = 0
        for student in students:
            student_fee, _ = StudentFee.objects.get_or_create(
                student=student, 
                academic_year=active_year,
                defaults={'total_fee': amount}
            )
            student_fee.total_fee = amount
            student_fee.save()
            updated_count += 1
        
        selected_sections = Section.objects.filter(id__in=section_ids).values_list('name', flat=True)
        sections_str = ", ".join(selected_sections)
        messages.success(request, f"Tuition fee ₹{amount} set for {updated_count} students across sections: {sections_str}.")
        return redirect('fee_list')

    return render(request, 'fees/bulk_tuition.html', {
        'sections': sections,
        'active_year': active_year,
    })


@all_roles_required
def fee_detail(request, pk):
    student = get_object_or_404(Student, pk=pk)
    active_year = AcademicYear.objects.filter(is_active=True).first()
    student_fee, _ = StudentFee.objects.get_or_create(
        student=student, academic_year=active_year, defaults={'total_fee': 0}
    )
    
    # Tuition payments
    tuition_payments = list(student_fee.payments.order_by('-payment_date', '-created_at'))
    
    # Other fee charges and payments
    other_charges = StudentFeeCharge.objects.filter(student=student, fee_type__academic_year=active_year).select_related('fee_type')
    other_payments = FeePayment.objects.filter(fee_charge__in=other_charges).order_by('-payment_date', '-created_at')

    all_payments = sorted(tuition_payments + list(other_payments), key=lambda p: (p.payment_date, p.created_at), reverse=True)

    return render(request, 'fees/detail.html', {
        'student': student, 'student_fee': student_fee,
        'other_charges': other_charges,
        'payments': all_payments, 'active_year': active_year,
    })


@admin_accounts_required
def fee_collect(request, pk):
    student = get_object_or_404(Student, pk=pk)
    active_year = AcademicYear.objects.filter(is_active=True).first()
    student_fee, _ = StudentFee.objects.get_or_create(
        student=student, academic_year=active_year, defaults={'total_fee': 0}
    )
    all_other_charges = StudentFeeCharge.objects.filter(student=student, fee_type__academic_year=active_year).select_related('fee_type').prefetch_related('payments')
    # Hide fully-paid charges from collect fee context
    unpaid_other_charges = [c for c in all_other_charges if c.total_pending > 0]

    today = timezone.localdate()

    if request.method == 'POST':
        fee_head_id = request.POST.get('fee_head', 'tuition')
        raw_amount = request.POST.get('amount', '0').strip()
        try:
            amount = float(raw_amount) if raw_amount else 0
        except ValueError:
            amount = 0
        payment_mode = request.POST.get('payment_mode', 'cash')
        remarks = request.POST.get('remarks', '')
        payment_date_str = request.POST.get('payment_date', str(today))
        custom_receipt = request.POST.get('receipt_number', '').strip()

        if amount <= 0:
            messages.error(request, 'Amount must be greater than 0.')
            return redirect('fee_collect', pk=pk)

        try:
            from datetime import date as Date
            payment_date = Date.fromisoformat(payment_date_str)
        except Exception:
            payment_date = today

        if custom_receipt:
            if FeePayment.objects.filter(receipt_number=custom_receipt).exists():
                messages.error(request, f'Receipt number {custom_receipt} already exists.')
                return redirect('fee_collect', pk=pk)
            receipt_number = custom_receipt
        else:
            receipt_number = 'RCP' + str(uuid.uuid4())[:8].upper()

        target_student_fee = None
        target_fee_charge = None

        if fee_head_id == 'tuition':
            if student_fee.total_fee == 0:
                messages.warning(request, f'Please set total tuition fee for {student.name} first.')
                return redirect('fee_set', pk=pk)
            target_student_fee = student_fee
        else:
            target_fee_charge = get_object_or_404(StudentFeeCharge, pk=fee_head_id, student=student)

        FeePayment.objects.create(
            student_fee=target_student_fee,
            fee_charge=target_fee_charge,
            amount=amount,
            payment_date=payment_date,
            payment_mode=payment_mode,
            receipt_number=receipt_number,
            collected_by=request.user.get_full_name() or request.user.username,
            remarks=remarks,
        )
        messages.success(request, f'₹{amount} collected! Receipt: {receipt_number}')
        next_url = request.POST.get('next', '').strip()
        if next_url:
            return redirect(next_url)
        return redirect('fee_detail', pk=pk)

    return render(request, 'fees/collect.html', {
        'student': student, 'student_fee': student_fee,
        'other_charges': unpaid_other_charges,
        'active_year': active_year, 'payment_modes': FeePayment.PAYMENT_MODES,
        'today': today,
    })


@admin_accounts_required
def payment_edit(request, pk, payment_id):
    student = get_object_or_404(Student, pk=pk)
    payment = get_object_or_404(FeePayment, pk=payment_id)

    if not (payment.student_fee and payment.student_fee.student == student) and not (payment.fee_charge and payment.fee_charge.student == student):
        messages.error(request, 'Payment record does not belong to this student.')
        return redirect('fee_detail', pk=pk)

    if request.method == 'POST':
        raw_amount = request.POST.get('amount', '').strip()
        payment_mode = request.POST.get('payment_mode', payment.payment_mode)
        payment_date_str = request.POST.get('payment_date', '')
        remarks = request.POST.get('remarks', '')

        new_receipt_number = request.POST.get('receipt_number', '').strip()
        if new_receipt_number and new_receipt_number != payment.receipt_number:
            if FeePayment.objects.filter(receipt_number=new_receipt_number).exclude(pk=payment.pk).exists():
                messages.error(request, f'Receipt number "{new_receipt_number}" is already used by another payment.')
                return redirect('payment_edit', pk=pk, payment_id=payment_id)
            payment.receipt_number = new_receipt_number

        try:
            amount = float(raw_amount)
            payment.amount = amount
        except ValueError:
            messages.error(request, 'Invalid amount value.')
            return redirect('payment_edit', pk=pk, payment_id=payment_id)

        if payment_date_str:
            try:
                from datetime import date as Date
                payment.payment_date = Date.fromisoformat(payment_date_str)
            except ValueError:
                pass

        payment.payment_mode = payment_mode
        payment.remarks = remarks
        payment.save()

        messages.success(request, f'Payment receipt #{payment.receipt_number} updated successfully.')
        return redirect('fee_detail', pk=pk)

    return render(request, 'fees/payment_edit.html', {
        'student': student,
        'payment': payment,
        'payment_modes': FeePayment.PAYMENT_MODES,
    })


@admin_accounts_required
def payment_delete(request, pk, payment_id):
    student = get_object_or_404(Student, pk=pk)
    payment = get_object_or_404(FeePayment, pk=payment_id)

    if not (payment.student_fee and payment.student_fee.student == student) and not (payment.fee_charge and payment.fee_charge.student == student):
        messages.error(request, 'Payment record does not belong to this student.')
        return redirect('fee_detail', pk=pk)

    if request.method == 'POST':
        receipt_no = payment.receipt_number
        amount = payment.amount
        payment.delete()
        messages.success(request, f'Payment receipt #{receipt_no} (₹{amount}) deleted successfully.')
        return redirect('fee_detail', pk=pk)

    return render(request, 'fees/payment_delete_confirm.html', {
        'student': student,
        'payment': payment,
    })


@admin_accounts_required
def fee_charge_set(request, pk, charge_id):
    student = get_object_or_404(Student, pk=pk)
    charge = get_object_or_404(StudentFeeCharge, pk=charge_id, student=student)

    if request.method == 'POST':
        raw_amount = request.POST.get('amount_assigned', '0').strip()
        try:
            charge.amount_assigned = float(raw_amount) if raw_amount else 0
        except ValueError:
            charge.amount_assigned = 0
        charge.save()
        messages.success(request, f"Assigned fee for {charge.fee_type.name} set to ₹{charge.amount_assigned} for {student.name}.")
        return redirect('fee_detail', pk=pk)

    return render(request, 'fees/set_charge.html', {
        'student': student,
        'charge': charge,
    })


@admin_accounts_required
def payment_adjust(request, pk):
    student = get_object_or_404(Student, pk=pk)
    active_year = AcademicYear.objects.filter(is_active=True).first()
    student_fee, _ = StudentFee.objects.get_or_create(
        student=student, academic_year=active_year, defaults={'total_fee': 0}
    )
    other_charges = StudentFeeCharge.objects.filter(student=student, fee_type__academic_year=active_year).select_related('fee_type')

    if request.method == 'POST':
        fee_head_id = request.POST.get('fee_head', 'tuition')
        raw_amount = request.POST.get('amount', '0').strip()
        try:
            amount = float(raw_amount) if raw_amount else 0
        except ValueError:
            amount = 0

        remarks = request.POST.get('remarks', 'Manual correction by admin').strip()
        if not remarks:
            remarks = 'Manual correction by admin'

        payment_date_str = request.POST.get('payment_date', str(timezone.localdate()))
        try:
            from datetime import date as Date
            payment_date = Date.fromisoformat(payment_date_str)
        except Exception:
            payment_date = timezone.localdate()

        receipt_number = 'ADJ' + str(uuid.uuid4())[:8].upper()

        target_student_fee = None
        target_fee_charge = None
        if fee_head_id == 'tuition':
            target_student_fee = student_fee
        else:
            target_fee_charge = get_object_or_404(StudentFeeCharge, pk=fee_head_id, student=student)

        FeePayment.objects.create(
            student_fee=target_student_fee,
            fee_charge=target_fee_charge,
            amount=amount,
            payment_date=payment_date,
            payment_mode='cash',
            receipt_number=receipt_number,
            collected_by=request.user.get_full_name() or request.user.username,
            remarks=remarks,
        )
        messages.success(request, f'Fee adjustment of ₹{amount} saved! Receipt: {receipt_number}')
        return redirect('fee_detail', pk=pk)

    return render(request, 'fees/payment_adjust.html', {
        'student': student,
        'student_fee': student_fee,
        'other_charges': other_charges,
        'active_year': active_year,
        'today': timezone.localdate(),
    })


@all_roles_required
def receipt_download(request, pk, payment_id):
    """Generate and download/view the fee receipt PDF for a payment."""
    student = get_object_or_404(Student, pk=pk)
    payment = get_object_or_404(FeePayment, pk=payment_id)
    if not (payment.student_fee and payment.student_fee.student == student) and not (payment.fee_charge and payment.fee_charge.student == student):
        return HttpResponse("Unauthorized", status=401)

    pdf_bytes = generate_receipt_pdf(payment)
    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="receipt_{payment.receipt_number}.pdf"'
    return response


@admin_accounts_required
def receipt_send_whatsapp(request, pk, payment_id):
    """Generate the fee receipt PDF, upload it, and send it to the student's
    registered mobile number via WhatsApp."""
    student = get_object_or_404(Student, pk=pk)
    payment = get_object_or_404(FeePayment, pk=payment_id)
    if not (payment.student_fee and payment.student_fee.student == student) and not (payment.fee_charge and payment.fee_charge.student == student):
        return HttpResponse("Unauthorized", status=401)

    if not student.mobile:
        messages.error(request, f'{student.name} has no mobile number on file.')
        return redirect('fee_detail', pk=pk)

    try:
        relative_path, _ = save_receipt_pdf(payment)
    except Exception as e:
        messages.error(request, f'Could not generate/save receipt: {e}')
        return redirect('fee_detail', pk=pk)

    media_url = f"{settings.SITE_BASE_URL.rstrip('/')}{settings.MEDIA_URL}{relative_path}"

    caption = (
        f"Hi {student.father_name}, this is the fee receipt for {student.name} "
        f"(Receipt No: {payment.receipt_number}, Amount: Rs. {payment.amount}). "
        f"- Sri NRI Junior College"
    )

    result = send_whatsapp_media(
        to_number=student.mobile,
        media_url=media_url,
        caption=caption,
    )

    if result.get('success'):
        messages.success(request, f'Receipt sent to {student.name} ({student.mobile}) via WhatsApp.')
    else:
        messages.error(request, f'Failed to send WhatsApp message: {result.get("error")}')

    return redirect('fee_detail', pk=pk)


@all_roles_required
def fee_export(request):
    """Export fees filtered by section and year to PDF or Excel.
    Columns: Name, Section, Year, Total, Paid, Pending, Parent Phone
    """
    from accounts.models import Section as Sec, Group
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    section_id = request.GET.get('section', '')
    year_filter = request.GET.get('year', '')
    fmt = request.GET.get('fmt', 'excel')
    active_year = AcademicYear.objects.filter(is_active=True).first()

    students = Student.objects.filter(is_active=True).select_related('section__group')
    if section_id:
        students = students.filter(section_id=section_id)
    if year_filter:
        students = students.filter(section__year=year_filter)

    fee_rows = []
    for student in students:
        sf = StudentFee.objects.filter(student=student, academic_year=active_year).first()
        fee_rows.append({
            'name':    student.name,
            'section': str(student.section) if student.section else '—',
            'year':    student.section.get_year_display() if student.section else '—',
            'total':   float(sf.total_fee) if sf else 0,
            'paid':    float(sf.total_paid) if sf else 0,
            'pending': float(sf.total_pending) if sf else 0,
            'phone':   student.mobile or '—',
        })

    if fmt == 'pdf':
        return _fee_export_pdf(fee_rows, active_year)
    return _fee_export_excel(fee_rows, active_year)


def _fee_export_excel(fee_rows, active_year):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fee Report"

    BLUE  = PatternFill("solid", fgColor="1A47A8")
    GREEN = PatternFill("solid", fgColor="DCFCE7")
    RED   = PatternFill("solid", fgColor="FEE2E2")
    LGREY = PatternFill("solid", fgColor="EAF1FB")
    thin  = Border(left=Side(style='thin'), right=Side(style='thin'),
                   top=Side(style='thin'), bottom=Side(style='thin'))

    # Title
    ws.merge_cells('A1:G1')
    ws['A1'] = f"Sri NRI Junior College — Fee Report ({active_year})"
    ws['A1'].font = Font(bold=True, size=13, color="1A47A8")
    ws['A1'].alignment = Alignment(horizontal='center')

    headers = ['S.No', 'Student Name', 'Section', 'Year', 'Total Fee', 'Paid', 'Pending', 'Parent Phone']
    col_widths = [6, 30, 20, 10, 14, 14, 14, 18]
    for i, (h, w) in enumerate(zip(headers, col_widths), 1):
        c = ws.cell(2, i, h)
        c.fill = BLUE
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal='center')
        c.border = thin
        ws.column_dimensions[c.column_letter].width = w

    for idx, row in enumerate(fee_rows, 1):
        r = idx + 2
        bg = LGREY if idx % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")
        vals = [idx, row['name'], row['section'], row['year'],
                row['total'], row['paid'], row['pending'], row['phone']]
        for col, val in enumerate(vals, 1):
            c = ws.cell(r, col, val)
            c.border = thin
            c.alignment = Alignment(horizontal='center' if col != 2 else 'left')
            if col == 6:   c.fill = GREEN
            elif col == 7: c.fill = RED
            else:          c.fill = bg

    # Totals row
    r = len(fee_rows) + 3
    ws.cell(r, 1, 'TOTAL').font = Font(bold=True)
    ws.cell(r, 5, sum(x['total'] for x in fee_rows)).font = Font(bold=True)
    ws.cell(r, 6, sum(x['paid'] for x in fee_rows)).font = Font(bold=True)
    ws.cell(r, 7, sum(x['pending'] for x in fee_rows)).font = Font(bold=True)

    from django.http import HttpResponse
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=fee_report.xlsx'
    wb.save(response)
    return response


def _fee_export_pdf(fee_rows, active_year):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import ParagraphStyle
    from django.http import HttpResponse
    import io

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=12*mm, rightMargin=12*mm,
                            topMargin=12*mm, bottomMargin=12*mm)
    DARK_BLUE = colors.HexColor("#1A47A8")
    GREEN     = colors.HexColor("#DCFCE7")
    RED       = colors.HexColor("#FEE2E2")
    LGREY     = colors.HexColor("#EAF1FB")

    story = []
    story.append(Paragraph(
        f"Sri NRI Junior College — Fee Report ({active_year})",
        ParagraphStyle('t', fontSize=13, fontName='Helvetica-Bold',
                       textColor=DARK_BLUE, alignment=1, spaceAfter=8)
    ))

    headers = ['#', 'Name', 'Section', 'Yr', 'Total', 'Paid', 'Pending', 'Phone']
    data = [headers]
    for i, row in enumerate(fee_rows, 1):
        data.append([
            str(i), row['name'], row['section'], row['year'],
            f"₹{row['total']:,.0f}", f"₹{row['paid']:,.0f}",
            f"₹{row['pending']:,.0f}", row['phone'],
        ])
    # Totals
    data.append(['', 'TOTAL', '', '',
                 f"₹{sum(x['total'] for x in fee_rows):,.0f}",
                 f"₹{sum(x['paid'] for x in fee_rows):,.0f}",
                 f"₹{sum(x['pending'] for x in fee_rows):,.0f}", ''])

    col_w = [8*mm, 45*mm, 28*mm, 12*mm, 22*mm, 22*mm, 22*mm, 28*mm]
    t = Table(data, colWidths=col_w, repeatRows=1)
    n = len(fee_rows)
    style = [
        ('BACKGROUND', (0,0), (-1,0), DARK_BLUE),
        ('TEXTCOLOR',  (0,0), (-1,0), colors.white),
        ('FONTNAME',   (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE',   (0,0), (-1,-1), 7.5),
        ('ROWBACKGROUNDS', (0,1), (-1,n), [colors.white, LGREY]),
        ('BACKGROUND', (5,1), (5,n), GREEN),
        ('BACKGROUND', (6,1), (6,n), RED),
        ('BACKGROUND', (0,n+1), (-1,n+1), colors.HexColor("#EAF1FB")),
        ('FONTNAME',   (0,n+1), (-1,n+1), 'Helvetica-Bold'),
        ('ALIGN',      (0,0), (-1,-1), 'CENTER'),
        ('ALIGN',      (1,0), (1,-1), 'LEFT'),
        ('GRID',       (0,0), (-1,-1), 0.3, colors.HexColor("#D1D5DB")),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]
    t.setStyle(TableStyle(style))
    story.append(t)
    doc.build(story)
    buf.seek(0)
    response = HttpResponse(buf, content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename=fee_report.pdf'
    return response
