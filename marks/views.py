from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.http import HttpResponse
from django.conf import settings
from .models import Exam, Mark, ExamSubjectMaxMark, ExamCategory, ExamType, GroupCategoryConfig, Subject, MarksEntryLock, MarksWhatsAppSendLog
from students.models import Student
from accounts.models import Section, AcademicYear, Group
from accounts.decorators import all_roles_required, admin_faculty_required, admin_required

def get_subjects_for_exam(exam):
    """Return the subject list to use for marks entry/report for this exam."""
    if exam and exam.category_id:
        all_subs = list(exam.category.subjects.all())
        if hasattr(exam, 'excluded_subjects'):
            excluded_ids = set(exam.excluded_subjects.values_list('id', flat=True))
            if excluded_ids:
                return [s for s in all_subs if s.id not in excluded_ids]
        return all_subs
    return []

def filter_subjects_for_user(request, subjects):
    """Filter subject list based on faculty assigned_subjects if user is faculty."""
    if getattr(request.user, 'role', None) == 'faculty' and not request.user.is_superuser:
        faculty_profile = getattr(request.user, 'faculty_profile', None)
        if faculty_profile:
            assigned_subs = set(faculty_profile.assigned_subjects.all())
            return [s for s in subjects if s in assigned_subs]
        return []
    return subjects

def get_subject_max_marks(exam, subjects):
    """Return {subject: max_marks} for the given exam."""
    saved = {m.subject: m.max_marks for m in exam.subject_max_marks.select_related('subject')}
    return {sub: saved.get(sub, exam.max_marks) for sub in subjects}

def get_section_completion_status(exam):
    """Return section completion status list for every Section under exam.group."""
    subjects = get_subjects_for_exam(exam)
    total_subjects_count = len(subjects)
    if exam.group_id:
        sections = list(Section.objects.filter(group=exam.group).select_related('group'))
    else:
        sections = list(Section.objects.all().select_related('group'))

    locks = MarksEntryLock.objects.filter(exam=exam, subject__in=subjects, is_locked=True)
    section_locks_map = {}
    for l in locks:
        section_locks_map.setdefault(l.section_id, set()).add(l.subject_id)

    sent_logs = {log.section_id: log for log in MarksWhatsAppSendLog.objects.filter(exam=exam)}

    result = []
    for sec in sections:
        locked_subs = section_locks_map.get(sec.id, set())
        completed_count = len(locked_subs)
        is_complete = (total_subjects_count > 0 and completed_count == total_subjects_count)
        send_log = sent_logs.get(sec.id)
        is_sent = bool(send_log)

        result.append({
            'section': sec,
            'is_complete': is_complete,
            'completed_subject_count': completed_count,
            'total_subject_count': total_subjects_count,
            'is_sent': is_sent,
            'send_log': send_log,
        })
    return result

@all_roles_required
def exam_list(request):
    active_year = AcademicYear.objects.filter(is_active=True).first()
    all_exams = (
        Exam.objects.filter(academic_year=active_year).select_related('group', 'category', 'exam_type').order_by('-date')
        if active_year else []
    )
    categories_qs = ExamCategory.objects.all()
    exam_boxes = []
    for cat in categories_qs:
        cat_exams = [e for e in all_exams if e.category_id == cat.id]
        exam_boxes.append({
            'key': str(cat.id), 'label': cat.name,
            'bg': cat.bg_color, 'color': cat.text_color, 'icon': cat.icon,
            'exams': cat_exams,
        })
    return render(request, 'marks/exam_list.html', {
        'exams': all_exams,
        'exam_boxes': exam_boxes,
        'active_year': active_year,
    })

@admin_faculty_required
def exam_add(request):
    active_year = AcademicYear.objects.filter(is_active=True).first()
    groups = Group.objects.all().order_by('name')
    exam_types = ExamType.objects.all()
    categories_qs = ExamCategory.objects.all()
    
    cl_data = {str(c.id): c.name for c in categories_qs}
    fc_data = [str(c.id) for c in categories_qs if c.is_fixed_marks]
    
    gco_data = {}
    ss_data = {}
    for config in GroupCategoryConfig.objects.select_related('group', 'category'):
        g_pk = str(config.group.pk)
        c_pk = str(config.category.pk)
        gco_data.setdefault(g_pk, []).append(c_pk)

    for cat in categories_qs.prefetch_related('subjects'):
        ss_data[str(cat.id)] = [{'id': sub.id, 'name': sub.name} for sub in cat.subjects.all()]
        
    exam_types_js = [[str(et.id), et.name] for et in exam_types]

    if request.method == 'POST':
        exam_type_id = request.POST.get('exam_type')
        custom_name = request.POST.get('custom_name', '').strip()
        group_id = request.POST.get('group_id')
        category_id = request.POST.get('category')

        if not group_id:
            messages.error(request, 'Please select a Student Group.')
        elif not category_id:
            messages.error(request, 'Please select an Exam Category.')
        else:
            group = get_object_or_404(Group, pk=group_id)
            category = get_object_or_404(ExamCategory, pk=category_id)
            exam_type = get_object_or_404(ExamType, pk=exam_type_id)

            if not GroupCategoryConfig.objects.filter(group=group, category=category).exists():
                messages.error(request, f'Category "{category.name}" is not configured for group "{group.name}".')
            else:
                subjects = list(category.subjects.all())
                if not subjects:
                    messages.error(request, f'Category "{category.name}" has no subjects attached. Please configure subjects in Admin first.')
                else:
                    subject_max_inputs = {}
                    if category.is_fixed_marks:
                        overall_max = 100 * len(subjects) # simple default for now, can be adjusted
                    else:
                        for subject in subjects:
                            val = request.POST.get(f'subject_max_{subject.name}', '').strip()
                            try:
                                subject_max_inputs[subject] = int(val) if val else 100
                            except (TypeError, ValueError):
                                subject_max_inputs[subject] = 100
                        overall_max = sum(subject_max_inputs.values()) or 100

                    exam = Exam.objects.create(
                        exam_type=exam_type,
                        custom_name=custom_name,
                        academic_year=active_year,
                        group=group,
                        category=category,
                        date=request.POST.get('date'),
                        max_marks=overall_max,
                    )

                    excluded_sub_ids = request.POST.getlist('excluded_subjects')
                    if excluded_sub_ids:
                        exam.excluded_subjects.set(excluded_sub_ids)

                    if not category.is_fixed_marks:
                        for subject, sub_max in subject_max_inputs.items():
                            ExamSubjectMaxMark.objects.update_or_create(
                                exam=exam, subject=subject, defaults={'max_marks': sub_max}
                            )

                    messages.success(request, 'Exam created successfully.')
                    return redirect('exam_list')

    return render(request, 'marks/exam_form.html', {
        'active_year': active_year, 'exam_types': exam_types,
        'groups': groups,
        'group_category_options': gco_data,
        'category_labels': cl_data,
        'subject_sets': ss_data,
        'fixed_categories': fc_data,
        'exam_types_js': exam_types_js,
        'default_max_marks': {str(c.id): 100 for c in categories_qs},
        'fixed_subject_max_marks': {str(c.id): {} for c in categories_qs if c.is_fixed_marks},
    })


@admin_faculty_required
def exam_edit(request, exam_id):
    active_year = AcademicYear.objects.filter(is_active=True).first()
    exam = get_object_or_404(Exam, pk=exam_id)
    all_category_subjects = list(exam.category.subjects.all())

    if request.method == 'POST':
        custom_name = request.POST.get('custom_name', '').strip()
        exam_date = request.POST.get('date')

        if custom_name is not None:
            exam.custom_name = custom_name
        if exam_date:
            exam.date = exam_date

        excluded_sub_ids = [int(i) for i in request.POST.getlist('excluded_subjects') if str(i).isdigit()]
        exam.excluded_subjects.set(excluded_sub_ids)
        exam.save()

        # Update max marks for included subjects
        active_subjects = [s for s in all_category_subjects if s.id not in set(excluded_sub_ids)]
        if not exam.category.is_fixed_marks and active_subjects:
            subject_max_inputs = {}
            for subject in active_subjects:
                val = request.POST.get(f'subject_max_{subject.name}', '').strip()
                try:
                    subject_max_inputs[subject] = int(val) if val else 100
                except (TypeError, ValueError):
                    subject_max_inputs[subject] = 100

                ExamSubjectMaxMark.objects.update_or_create(
                    exam=exam, subject=subject, defaults={'max_marks': subject_max_inputs[subject]}
                )
            exam.max_marks = sum(subject_max_inputs.values()) or 100
            exam.save()

        messages.success(request, f"Exam '{exam.display_name()}' updated successfully.")
        return redirect('exam_list')

    existing_max_marks = {
        m.subject_id: m.max_marks
        for m in ExamSubjectMaxMark.objects.filter(exam=exam)
    }
    excluded_subject_ids = set(exam.excluded_subjects.values_list('id', flat=True))

    return render(request, 'marks/exam_edit.html', {
        'exam': exam,
        'active_year': active_year,
        'subjects': all_category_subjects,
        'excluded_subject_ids': excluded_subject_ids,
        'existing_max_marks': existing_max_marks,
    })


@admin_required
def exam_delete(request, exam_id):
    exam = get_object_or_404(Exam, pk=exam_id)
    name = exam.display_name()
    exam.delete()
    messages.success(request, f"Exam '{name}' and all associated marks have been deleted.")
    return redirect('exam_list')


@admin_required
def marks_entry_unlock(request, exam_id, section_id, subject_id):
    exam = get_object_or_404(Exam, pk=exam_id)
    section = get_object_or_404(Section, pk=section_id)
    subject = get_object_or_404(Subject, pk=subject_id)
    MarksEntryLock.objects.update_or_create(
        exam=exam, section=section, subject=subject,
        defaults={'is_locked': False, 'locked_by': request.user}
    )
    messages.success(request, f'Unlocked {subject.name} for editing.')
    return redirect(f'/marks/exam/{exam_id}/entry/?section={section_id}')

@admin_faculty_required
def marks_entry(request, exam_id):
    exam = get_object_or_404(Exam, pk=exam_id)
    section_id = request.GET.get('section', '')
    if exam.group_id:
        sections = Section.objects.select_related('group').filter(group=exam.group)
    else:
        sections = Section.objects.select_related('group').all()

    subjects = get_subjects_for_exam(exam)
    subjects = filter_subjects_for_user(request, subjects)

    if not subjects:
        if getattr(request.user, 'role', None) == 'faculty':
            messages.warning(request, 'You are not assigned to enter marks for any subject in this exam.')
        else:
            messages.warning(request, f'No subjects are attached to exam category "{exam.category.name}". Please attach subjects in Django Admin.')

    subject_max_marks = get_subject_max_marks(exam, subjects)
    rows = []
    selected_section = None
    locked_map = {}

    if section_id and str(section_id).isdigit():
        selected_section = get_object_or_404(Section, pk=section_id)
        locks = MarksEntryLock.objects.filter(exam=exam, section=selected_section, subject__in=subjects)
        locked_subject_ids = set(locks.filter(is_locked=True).values_list('subject_id', flat=True))
        locked_map = {sub.id: (sub.id in locked_subject_ids) for sub in subjects}

    if request.method == 'POST':
        sid = request.POST.get('section_id')
        sec = get_object_or_404(Section, pk=sid)

        # Server-side lock check
        locks = MarksEntryLock.objects.filter(exam=exam, section=sec, subject__in=subjects)
        locked_subject_ids = set(locks.filter(is_locked=True).values_list('subject_id', flat=True))
        is_faculty = (getattr(request.user, 'role', None) == 'faculty' and not request.user.is_superuser)
        saved_subjects_set = set()

        for student in Student.objects.filter(section=sec, is_active=True):
            for subject in subjects:
                if is_faculty and subject.id in locked_subject_ids:
                    continue

                val = request.POST.get(f'mark_{student.pk}_{subject.name}', '').strip()
                if val:
                    try:
                        num_val = float(val)
                        is_absent_val = (num_val == 0)
                    except (ValueError, TypeError):
                        is_absent_val = False

                    Mark.objects.update_or_create(
                        student=student, exam=exam, subject=subject,
                        defaults={'marks_obtained': val, 'is_absent': is_absent_val}
                    )
                    saved_subjects_set.add(subject)

        # Lock saved subjects for both Admin and Faculty
        for subject in saved_subjects_set:
            MarksEntryLock.objects.update_or_create(
                exam=exam, section=sec, subject=subject,
                defaults={'is_locked': True, 'locked_by': request.user}
            )

        messages.success(request, 'Marks saved successfully.')
        return redirect(f'/marks/exam/{exam_id}/entry/?section={sid}')

    if selected_section:
        students = list(Student.objects.filter(section=selected_section, is_active=True))
        existing = Mark.objects.filter(exam=exam, student__in=students)
        marks_map = {}
        for m in existing:
            marks_map.setdefault(m.student_id, {})[m.subject_id] = m.marks_obtained or ''
        for student in students:
            s_marks = [(sub, marks_map.get(student.pk, {}).get(sub.id, '')) for sub in subjects]
            rows.append({'student': student, 'marks': s_marks})

    subjects_with_info = [
        (sub, subject_max_marks.get(sub, exam.max_marks), locked_map.get(sub.id, False))
        for sub in subjects
    ]

    return render(request, 'marks/entry.html', {
        'exam': exam, 'sections': sections, 'rows': rows,
        'selected_section': selected_section, 'section_id': str(section_id),
        'subjects': subjects, 'subject_max_marks': subject_max_marks,
        'subjects_with_info': subjects_with_info,
        'locked_map': locked_map,
    })

@admin_required
def marks_whatsapp_send(request, exam_id):
    from whatsapp.models import PendingMessage
    exam = get_object_or_404(Exam, pk=exam_id)
    statuses = get_section_completion_status(exam)

    total_sections = len(statuses)
    sent_sections_count = sum(1 for s in statuses if s['is_sent'])
    ready_sections_count = sum(1 for s in statuses if s['is_complete'])
    pending_marks_count = sum(1 for s in statuses if not s['is_complete'])

    if request.method == 'POST':
        selected_section_ids = request.POST.getlist('selected_sections')
        target_lang = (request.POST.get('language') or 'en').lower().strip()
        if not selected_section_ids:
            messages.warning(request, 'No section selected to send.')
            return redirect(f'/marks/exam/{exam_id}/whatsapp/send/')

        sec_map = {str(s['section'].id): s for s in statuses}
        target_statuses = []
        for sid in selected_section_ids:
            status = sec_map.get(str(sid))
            if status and status['is_complete']:
                target_statuses.append(status)

        if not target_statuses:
            messages.warning(request, 'Selected section(s) are incomplete. Please complete marks entry first.')
            return redirect(f'/marks/exam/{exam_id}/whatsapp/send/')

        subjects = get_subjects_for_exam(exam)
        subject_max_marks = get_subject_max_marks(exam, subjects)

        total_queued_messages = 0
        sent_section_names = []
        template_name = getattr(settings, 'META_TEMPLATE_EXAM_MARKS', 'exam_marks')

        for st in target_statuses:
            sec = st['section']
            students = Student.objects.filter(section=sec, is_active=True)
            if not students.exists():
                MarksWhatsAppSendLog.objects.create(
                    exam=exam, section=sec, sent_by=request.user, student_count=0
                )
                sent_section_names.append(str(sec))
                continue

            marks_qs = Mark.objects.filter(exam=exam, student__in=students).select_related('subject')
            student_marks_map = {}
            for m in marks_qs:
                student_marks_map.setdefault(m.student_id, {})[m.subject_id] = m

            sec_message_count = 0
            for student in students:
                phone = (student.mobile or getattr(student, 'second_mobile', '') or '').strip()
                if not phone:
                    continue

                s_marks = student_marks_map.get(student.id, {})
                mark_lines = []
                total_obtained = 0.0
                total_max = 0
                has_absent = False

                for sub in subjects:
                    m = s_marks.get(sub.id)
                    s_max = subject_max_marks.get(sub, exam.max_marks)
                    total_max += s_max
                    if m:
                        is_abs = m.is_absent or (m.marks_obtained is not None and float(m.marks_obtained) == 0)
                        if is_abs:
                            has_absent = True
                            mark_lines.append(f"• {sub.name}: AB / {s_max}")
                        elif m.marks_obtained is not None:
                            val = float(m.marks_obtained)
                            total_obtained += val
                            mark_lines.append(f"• {sub.name}: {val:g} / {s_max}")
                        else:
                            mark_lines.append(f"• {sub.name}: - / {s_max}")
                    else:
                        has_absent = True
                        mark_lines.append(f"• {sub.name}: AB / {s_max}")

                pct = round(total_obtained / total_max * 100, 1) if total_max else 0

                if has_absent:
                    obtained_str = "ABSENT"
                    exam_status_str = f"{exam.display_name()} (ABSENT)"
                    total_obtained_msg = f"Total Obtained: ABSENT / {total_max}"
                else:
                    obtained_str = f"{total_obtained:g}"
                    exam_status_str = f"{exam.display_name()} ({pct}%)"
                    total_obtained_msg = f"Total Obtained: {total_obtained:g} / {total_max} ({pct}%)"

                msg_text = (
                    f"Sri NRI Junior College — Marks Report\n"
                    f"Student: {student.name} ({student.admission_number})\n"
                    f"Exam: {exam.display_name()} ({exam.category.name})\n"
                    f"Date: {exam.date.strftime('%d-%m-%Y')}\n\n"
                    f"Subject-wise Marks:\n" + "\n".join(mark_lines) + "\n\n"
                    f"{total_obtained_msg}"
                )

                subject_summary = "\n".join(mark_lines)

                template_params = [
                    student.name,
                    obtained_str,
                    str(total_max),
                    exam_status_str,
                    exam.date.strftime('%d-%m-%Y'),
                ]
                if getattr(settings, 'MARKS_TEMPLATE_HAS_SUBJECTS', False):
                    template_params.append(subject_summary)

                PendingMessage.objects.create(
                    student=student,
                    phone=phone,
                    message_type=PendingMessage.TYPE_TEMPLATE,
                    template_name=template_name,
                    template_params=template_params,
                    language=target_lang,
                    message=msg_text,
                    status=PendingMessage.STATUS_PENDING,
                )
                sec_message_count += 1
                total_queued_messages += 1

            MarksWhatsAppSendLog.objects.update_or_create(
                exam=exam,
                section=sec,
                defaults={
                    'sent_by': request.user,
                    'student_count': sec_message_count,
                }
            )
            sent_section_names.append(str(sec))

        if total_queued_messages > 0:
            from whatsapp.services import dispatch_pending_messages_async
            dispatch_pending_messages_async()

        messages.success(
            request,
            f"Successfully queued {total_queued_messages} WhatsApp message(s) for section(s): {', '.join(sent_section_names)}."
        )
        return redirect(f'/marks/exam/{exam_id}/whatsapp/send/')

    return render(request, 'marks/whatsapp_send.html', {
        'exam': exam,
        'statuses': statuses,
        'total_sections': total_sections,
        'sent_sections_count': sent_sections_count,
        'ready_sections_count': ready_sections_count,
        'pending_marks_count': pending_marks_count,
    })

@all_roles_required
def marks_report(request):
    import json
    group_id   = request.GET.get('group', '')
    category   = request.GET.get('category', '')
    section_id = request.GET.get('section', '')
    exam_id    = request.GET.get('exam', '')
    sort_order = request.GET.get('sort', 'asc')

    groups   = Group.objects.all().order_by('name')
    sections = Section.objects.select_related('group').all()
    active_year = AcademicYear.objects.filter(is_active=True).first()
    exams = Exam.objects.filter(academic_year=active_year).select_related('group', 'category', 'exam_type').order_by('-date') if active_year else []
    categories_qs = ExamCategory.objects.all()

    group_category_map = {}
    for config in GroupCategoryConfig.objects.select_related('group', 'category').all():
        group_category_map.setdefault(str(config.group_id), []).append(str(config.category_id))
    group_category_map_json = json.dumps(group_category_map)

    report_data = []
    selected_section = None
    selected_exam = None
    section_label = ''
    subjects = []
    subject_max_marks = {}
    total_max = 0
    is_all_exams = False
    all_exams_list = []

    eff_group    = None if (not group_id    or group_id    == 'all') else group_id
    eff_category = None if (not category    or category    == 'all') else category
    eff_section  = None if (not section_id  or section_id  == 'all') else section_id

    def _resolve_students():
        nonlocal selected_section, section_label
        if eff_section:
            selected_section = get_object_or_404(Section, pk=eff_section)
            section_label = str(selected_section)
            return Student.objects.filter(section=selected_section, is_active=True)
        qs = Student.objects.select_related('section').filter(is_active=True)
        if eff_group:
            qs = qs.filter(section__group_id=eff_group)
        section_label = 'All Sections'
        return qs

    if section_id and exam_id:
        if exam_id == 'all':
            is_all_exams = True
            students = _resolve_students()

            exam_qs = Exam.objects.filter(academic_year=active_year).select_related('group').order_by('date')
            if eff_group:
                exam_qs = exam_qs.filter(group_id=eff_group)
            if eff_category:
                exam_qs = exam_qs.filter(category_id=eff_category)
            all_exams_list = list(exam_qs)

            all_marks = Mark.objects.filter(exam__in=all_exams_list, student__in=students)
            exam_total_map = {}
            for m in all_marks:
                is_abs = m.is_absent or (m.marks_obtained is not None and float(m.marks_obtained) == 0)
                if m.marks_obtained and not is_abs:
                    exam_total_map.setdefault(m.student_id, {})
                    exam_total_map[m.student_id][m.exam_id] = (
                        exam_total_map[m.student_id].get(m.exam_id, 0) + float(m.marks_obtained)
                    )
            for student in students:
                s_totals = exam_total_map.get(student.pk, {})
                per_exam = [round(s_totals.get(e.pk, 0), 1) for e in all_exams_list]
                overall = round(sum(per_exam), 1)
                report_data.append({'student': student, 'per_exam': per_exam, 'total': overall})

            reverse = (sort_order == 'desc')
            report_data.sort(key=lambda r: r['total'], reverse=reverse)

        else:
            selected_exam = get_object_or_404(Exam, pk=exam_id)
            students = _resolve_students()

            marks = Mark.objects.filter(exam=selected_exam, student__in=students).select_related('subject')
            subjects = get_subjects_for_exam(selected_exam)
            subjects = filter_subjects_for_user(request, subjects)
            subject_max_marks = get_subject_max_marks(selected_exam, subjects)
            total_max = sum(subject_max_marks.get(sub, 0) for sub in subjects)
            marks_map = {}
            for m in marks:
                marks_map.setdefault(m.student_id, {})[m.subject_id] = m
            for student in students:
                s_marks = marks_map.get(student.pk, {})
                mark_list = [s_marks.get(sub.id) for sub in subjects]
                total = sum(float(m.marks_obtained) for m in s_marks.values()
                           if m.marks_obtained and not m.is_absent and float(m.marks_obtained) != 0)
                pct = round(total / total_max * 100, 1) if total_max else 0
                report_data.append({'student': student, 'mark_list': mark_list, 'total': total, 'pct': pct})

            reverse = (sort_order == 'desc')
            report_data.sort(key=lambda r: r['total'], reverse=reverse)

    subjects_with_max = [(sub, subject_max_marks.get(sub, 0)) for sub in subjects]

    return render(request, 'marks/report.html', {
        'groups': groups,
        'category_choices': [(c.id, c.name) for c in categories_qs],
        'group_category_map_json': group_category_map_json,
        'sections': sections, 'exams': exams, 'report_data': report_data,
        'selected_section': selected_section, 'selected_exam': selected_exam,
        'section_label': section_label,
        'is_all_exams': is_all_exams, 'all_exams_list': all_exams_list,
        'group_id': str(group_id), 'category': category,
        'section_id': str(section_id), 'exam_id': str(exam_id), 'subjects': subjects,
        'subject_max_marks': subject_max_marks, 'total_max': total_max,
        'sort_order': sort_order, 'subjects_with_max': subjects_with_max,
    })

@all_roles_required
def marks_report_export_excel(request):
    import openpyxl

    section_id = request.GET.get('section', '')
    exam_id = request.GET.get('exam', '')
    sort_order = request.GET.get('sort', 'asc')

    if not (section_id and exam_id):
        messages.error(request, 'Please select a section and exam before exporting.')
        return redirect('marks_report')

    selected_section = get_object_or_404(Section, pk=section_id)
    selected_exam = get_object_or_404(Exam, pk=exam_id)
    students = Student.objects.filter(section=selected_section, is_active=True)
    marks = Mark.objects.filter(exam=selected_exam, student__in=students).select_related('subject')
    subjects = get_subjects_for_exam(selected_exam)
    subjects = filter_subjects_for_user(request, subjects)
    subject_max_marks = get_subject_max_marks(selected_exam, subjects)

    marks_map = {}
    for m in marks:
        marks_map.setdefault(m.student_id, {})[m.subject_id] = m

    rows = []
    for student in students:
        s_marks = marks_map.get(student.pk, {})
        total = sum(float(m.marks_obtained) for m in s_marks.values()
                   if m.marks_obtained and not m.is_absent and float(m.marks_obtained) != 0)
        rows.append({'student': student, 's_marks': s_marks, 'total': total})

    reverse = (sort_order == 'desc')
    rows.sort(key=lambda r: r['total'], reverse=reverse)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Marks Report"

    headers = ['#', 'Admission No', 'Name'] + [s.name for s in subjects] + ['Total', f'Total (out of {sum(subject_max_marks.values())})']
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)

    for idx, row in enumerate(rows, 1):
        ws.cell(idx + 1, 1, idx)
        ws.cell(idx + 1, 2, row['student'].admission_number)
        ws.cell(idx + 1, 3, row['student'].name)
        for j, sub in enumerate(subjects):
            m = row['s_marks'].get(sub.id)
            if m:
                is_abs = m.is_absent or (m.marks_obtained is not None and float(m.marks_obtained) == 0)
                ws.cell(idx + 1, 4 + j, 'AB' if is_abs else float(m.marks_obtained or 0))
            else:
                ws.cell(idx + 1, 4 + j, '-')
        col_total = 4 + len(subjects)
        ws.cell(idx + 1, col_total, row['total'])
        ws.cell(idx + 1, col_total + 1, f"{row['total']}/{sum(subject_max_marks.values())}")

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f"marks_report_{selected_exam.display_name()}_{selected_section}.xlsx".replace(' ', '_')
    response['Content-Disposition'] = f'attachment; filename={filename}'
    wb.save(response)
    return response

@all_roles_required
def marks_report_export_pdf(request):
    from .marks_pdf import generate_marks_report_pdf

    section_id = request.GET.get('section', '')
    exam_id = request.GET.get('exam', '')
    sort_order = request.GET.get('sort', 'asc')

    if not (section_id and exam_id):
        messages.error(request, 'Please select a section and exam before exporting.')
        return redirect('marks_report')

    selected_section = get_object_or_404(Section, pk=section_id)
    selected_exam = get_object_or_404(Exam, pk=exam_id)
    students = Student.objects.filter(section=selected_section, is_active=True)
    marks = Mark.objects.filter(exam=selected_exam, student__in=students).select_related('subject')
    subjects = get_subjects_for_exam(selected_exam)
    subjects = filter_subjects_for_user(request, subjects)
    subject_max_marks = get_subject_max_marks(selected_exam, subjects)

    marks_map = {}
    for m in marks:
        marks_map.setdefault(m.student_id, {})[m.subject_id] = m

    rows = []
    for student in students:
        s_marks = marks_map.get(student.pk, {})
        total = sum(float(m.marks_obtained) for m in s_marks.values()
                   if m.marks_obtained and not m.is_absent)
        rows.append({'student': student, 's_marks': s_marks, 'total': total})

    reverse = (sort_order == 'desc')
    rows.sort(key=lambda r: r['total'], reverse=reverse)

    pdf_bytes = generate_marks_report_pdf(selected_exam, selected_section, subjects, subject_max_marks, rows)

    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    filename = f"marks_report_{selected_exam.display_name()}_{selected_section}.pdf".replace(' ', '_')
    response['Content-Disposition'] = f'attachment; filename={filename}'
    return response

@all_roles_required
def student_marks_export_pdf(request, student_id):
    from .marks_pdf import generate_student_marks_pdf

    student = get_object_or_404(Student, pk=student_id)
    marks = Mark.objects.filter(student=student).select_related('exam', 'subject').order_by('-exam__date')

    exams_map = {}
    for m in marks:
        exams_map.setdefault(m.exam_id, {'exam': m.exam, 'marks': {}})
        exams_map[m.exam_id]['marks'][m.subject_id] = m

    exam_rows = []
    for entry in exams_map.values():
        exam = entry['exam']
        subjects = get_subjects_for_exam(exam)
        subject_max_marks = get_subject_max_marks(exam, subjects)
        row_marks = []
        for sub in subjects:
            m = entry['marks'].get(sub.id)
            row_marks.append({
                'subject': sub.name,
                'mark': m,
                'max': subject_max_marks.get(sub, exam.max_marks),
            })
        exam_rows.append({'exam': exam, 'subject_marks': row_marks})

    pdf_bytes = generate_student_marks_pdf(student, exam_rows)

    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    filename = f"marks_{student.admission_number}_{student.name}.pdf".replace(' ', '_')
    response['Content-Disposition'] = f'attachment; filename={filename}'
    return response

@all_roles_required
def marks_report_export(request):
    fmt = request.GET.get('export', 'pdf')
    section_id = request.GET.get('section', '')
    exam_id    = request.GET.get('exam', '')
    if not (section_id and exam_id):
        messages.error(request, 'Please select section and exam first.')
        return redirect('marks_report')
    from django.urls import reverse
    base = reverse('marks_report_export_excel') if fmt == 'excel' else reverse('marks_report_export_pdf')
    return redirect(f"{base}?section={section_id}&exam={exam_id}")
